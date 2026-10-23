"""Check public Git content and optionally build a privacy-safe sharing ZIP.

Uses only staged/tracked public files, never raw local data or the old history.
Run after staging intended changes. Exports refuse to replace an existing ZIP.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import subprocess
import sys
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook

from generate_demo_data import datasets, workbook_sheets

VERSION = Path(__file__).resolve().parents[1]


def git(root, *args):
    return subprocess.check_output(['git','--no-optional-locks','-C',str(root),*args],stderr=subprocess.DEVNULL)


def approved(path, prefix):
    local = path[len(prefix):] if path.startswith(prefix) else None
    if local is not None:
        if local in ['.gitignore','README.md','requirements.txt','run_matching.py','tools/sample_hashes.json']:
            return True
        parts = Path(local).parts
        if not parts:
            return False
        if parts[0] == 'data':
            return len(parts) == 2 and parts[1] in datasets()
        suffixes = {'saturday_match':{'.py'},'tests':{'.py'},'tools':{'.py'},
                    'config':{'.yaml'},'docs':{'.md'}}
        return len(parts) >= 2 and parts[0] in suffixes and Path(local).suffix in suffixes[parts[0]]
    return path in ['README.md','.gitignore','.gitattributes','version_0/README.md',
                    'version_0/requirements.txt','version_0/2024dataframe.py','.github/workflows/checks.yml']


def zip_parts(payload):
    with ZipFile(io.BytesIO(payload)) as z:
        # Reject duplicate ZIP entries rather than hiding the earlier content.
        if len(z.namelist()) != len(set(z.namelist())):
            raise ValueError('duplicate_workbook_parts')
        return {name:z.read(name) for name in z.namelist()}


def sample_hashes(payload):
    """Read the approved, staged byte manifest; its updates require code review.

    Hashes are fixed public release artifacts, not recalculated from whichever
    workbooks happen to be present. A coordinated change to hashes and samples
    is a new approval decision: review the generator and synthetic declarations.
    """
    def unique_keys(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ValueError('duplicate_sample_manifest_key')
            result[key] = value
        return result
    try:
        hashes = json.loads(payload.decode('utf-8-sig'), object_pairs_hook=unique_keys)
    except (UnicodeError, json.JSONDecodeError):
        raise ValueError('invalid_sample_manifest_json') from None
    if not isinstance(hashes, dict) or set(hashes) != set(datasets()):
        raise ValueError('sample_manifest_must_list_exact_six_examples')
    if any(not isinstance(value, str) or re.fullmatch(r'[0-9a-f]{64}', value) is None for value in hashes.values()):
        raise ValueError('invalid_sample_manifest_digest')
    return hashes


def verify_sample(payload, filename, approved_hash):
    """Check all bytes before opening a workbook, then verify declared content.

    A member-content comparison alone misses archive comments, ZIP metadata,
    and trailing bytes. Fixed hashes cover those surfaces and do not depend on
    the installed openpyxl writer's formatting/version.
    """
    if hashlib.sha256(payload).hexdigest() != approved_hash:
        raise ValueError('not_exact_approved_synthetic_workbook')
    zip_parts(payload)  # Also reject duplicate archive names explicitly.
    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    try:
        declared_sheets = workbook_sheets()[filename]
        if book.sheetnames != list(declared_sheets):
            raise ValueError('unexpected_synthetic_worksheets')
        for name, declared_rows in declared_sheets.items():
            rows = [list(row) for row in book[name].iter_rows(values_only=True)]
            if rows != declared_rows:
                raise ValueError('synthetic_declaration_mismatch')
        expected = {
            'creator': 'Saturday matching synthetic examples',
            'lastModifiedBy': 'Saturday matching synthetic examples',
            'title': 'Synthetic questionnaire example',
            'description': 'Wholly invented participants; no historical rows.',
        }
        if any(getattr(book.properties, field) != value for field, value in expected.items()):
            raise ValueError('unexpected_synthetic_authorship')
        for field in ('created', 'modified'):
            value = getattr(book.properties, field)
            if value is None or value.isoformat()[:19] != '2000-01-01T00:00:00':
                raise ValueError('unexpected_synthetic_timestamp')
    finally:
        book.close()


def inspect():
    root = Path(git(VERSION,'rev-parse','--show-toplevel').decode().strip())
    prefix_path = VERSION.relative_to(root).as_posix()
    prefix = '' if prefix_path == '.' else prefix_path + '/'
    paths = [s for s in git(root,'ls-files','-z').decode().split('\0') if s]
    if not paths:
        raise ValueError('No tracked files: stage intended public files before release checks.')
    failures = []
    contents = {}
    for path in paths:
        if not approved(path,prefix):
            failures.append((path,'unexpected_tracked_file'))
            continue
        staged = git(root,'show',':'+path)
        contents[path] = staged
    manifest_path = prefix + 'tools/sample_hashes.json'
    hashes = None
    if manifest_path not in contents:
        failures.append((manifest_path,'approved_sample_manifest_not_tracked'))
    else:
        hashes = sample_hashes(contents[manifest_path])
    for path, staged in contents.items():
        if path.endswith('.xlsx'):
            name = Path(path).name
            if hashes is not None:
                try:
                    verify_sample(staged,name,hashes[name])
                except (ValueError, BadZipFile, KeyError):
                    failures.append((path,'not_exact_approved_synthetic_workbook'))
        else:
            text = staged.decode('utf-8-sig')
            if re.search(r'[A-Za-z]:\\Users\\[^\\\r\n]+',text):
                failures.append((path,'personal_account_path'))
            if re.search(r'gh[pousr]_[A-Za-z0-9]{20,}|sk-[A-Za-z0-9_-]{24,}|AKIA[A-Z0-9]{16}|-----BEGIN (?:RSA |OPENSSH )?PRIVATE KEY-----',text):
                failures.append((path,'credential_pattern'))
            for domain in re.findall(r'[\w.+-]+@([\w.-]+\.[A-Za-z]{2,})',text):
                if not (domain.endswith('.invalid') or domain in ['example.com','example.org','example.net']):
                    failures.append((path,'non_example_email'))
                    break
    for filename in datasets():
        if prefix+'data/'+filename not in paths:
            failures.append((prefix+'data/'+filename,'synthetic_example_not_tracked'))
    # Prevent a passing staged check from hiding unreviewed working-tree edits.
    dirty = subprocess.run(['git','--no-optional-locks','-C',str(root),'diff','--quiet'],capture_output=True).returncode
    if dirty:
        failures.append(('<working-tree>','unstaged_changes'))
    if failures:
        for path, code in failures:
            print(f'{path}: {code}',file=sys.stderr)
        raise ValueError(f'Release checks failed: {len(failures)} issue(s).')
    return contents


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--export',type=Path,help='Create a new ZIP from verified public staged contents.')
    a=p.parse_args()
    try:
        contents=inspect()
        if a.export:
            # No private .git, archive or original workbooks can enter this list.
            with a.export.open('xb') as stream, ZipFile(stream,'w',ZIP_DEFLATED) as z:
                for path,payload in sorted(contents.items()):
                    z.writestr(path,payload)
        print(f'Release checks passed: {len(contents)} public tracked files; six exact synthetic workbooks.')
        return 0
    except (ValueError,OSError,BadZipFile,subprocess.CalledProcessError) as exc:
        # Never print contents, source values or full subprocess output.
        print(str(exc) if isinstance(exc,ValueError) else 'Release files/Git could not be read.',file=sys.stderr)
        return 1


if __name__=='__main__':
    raise SystemExit(main())
