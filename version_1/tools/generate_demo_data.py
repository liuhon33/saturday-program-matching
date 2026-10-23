"""Generate six small, wholly invented questionnaire examples.

No historical workbook is read. Existing files are never overwritten. The
fixed output location is version_1/data, regardless of the current directory.
Run with --check to compare delivered workbooks with the declared examples.
"""

from __future__ import annotations

import argparse
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import re
import sys
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DATA_DIR = Path(__file__).resolve().parents[1] / "data"
FIXED_TIME = datetime(2000, 1, 1)

TUTOR_HEADERS = [
    "Tutor ID", "Tutor First Name", "Tutor Last Name", "Tutor Subjects", "Tutor Gender",
]
TUTEE_HEADERS = [
    "Tutee ID", "Tutee 1 First Name", "Tutee 1 Last Name", "Tutee 1 Primary Subject",
    "Tutee 1 Secondary Subject", "Tutee 1 Tertiary Subject", "Gender Preference",
]


def datasets() -> dict[str, list[list[str | None]]]:
    """Return invented cases; only form headers/academic choices mirror the app."""
    tutor_choices = [
        ("English; Social Sciences", "Female"),
        ("Math Grades 7, 8, 9, 10; Functions (Grade 11)", "Male"),
        ("Biology; Chemistry", "Female"),
        ("Physics; Calculus and/or Advanced Functions (Grade 12)", "Male"),
        ("French; English", "Female"),
        ("Computer Science; Data Management (Grade 12)", "Male"),
        ("Science Grades 7, 8, 9, 10; Chemistry", "Female"),
        ("English; History", "Male"),
    ]
    tutee_choices = [
        ("English", "French", None, "No preference"),
        ("Math Grades 7, 8, 9, 10", "Functions (Grade 11)", None, "No preference"),
        ("Biology", "Chemistry", None, "No preference"),
        ("Physics", "Calculus and/or Advanced Functions (Grade 12)", None, "Male"),
        ("French", "English", None, "Female"),
        ("Computer Science", "Data Management (Grade 12)", None, "No preference"),
        ("Science Grades 7, 8, 9, 10", "Chemistry", "Biology", "No preference"),
        ("History", "Social Sciences", "English", "No preference"),
        ("English", "Social Sciences", None, "Female"),
        ("Math Grades 7, 8, 9, 10", "Physics", None, "No preference"),
        ("Chemistry", "Biology", None, "No preference"),
        ("Calculus and/or Advanced Functions (Grade 12)", "Data Management (Grade 12)", None, "No preference"),
    ]
    result: dict[str, list[list[str | None]]] = {}
    for case, prefix in (("2023", "2023"), ("2024", "2024"), ("scratch", "scratch_")):
        tutors: list[list[str | None]] = [TUTOR_HEADERS.copy()]
        tutees: list[list[str | None]] = [TUTEE_HEADERS.copy()]
        for number, (subjects, gender) in enumerate(tutor_choices, start=1):
            # IDs 001 and 008 intentionally share a full display name.
            name_number = 1 if number == 8 else number
            tutors.append([f"demo-{case}-tutor-{number:03d}", "tutor", f"{name_number:03d}", subjects, gender])
        for number, (primary, secondary, tertiary, preference) in enumerate(tutee_choices, start=1):
            # IDs 001 and 012 intentionally share a full display name.
            name_number = 1 if number == 12 else number
            tutees.append([f"demo-{case}-tutee-{number:03d}", "tutee", f"{name_number:03d}", primary, secondary, tertiary, preference])
        if case == "2023":
            tutors[2][3] = "Math Grades 7, 8, 9, 10; Math Grade 11"
            tutees[2][4] = "Grade 11 Math"
        if case == "scratch":
            tutors[1][3] = "  english ; Social Sciences  "
            tutors[1][4] = " female "
            tutees[1][3] = "Englsh"  # Explicitly approved demo typo in the config.
            tutees[2][6] = " no preference "
        result[f"{prefix}tutor.xlsx"] = tutors
        result[f"{prefix}tutee.xlsx"] = tutees
    return result


def workbook_bytes(rows: list[list[str | None]]) -> bytes:
    """Make a deterministic workbook with neutral authorship and fixed dates."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    workbook.properties.creator = "Saturday matching synthetic examples"
    workbook.properties.lastModifiedBy = "Saturday matching synthetic examples"
    workbook.properties.title = "Synthetic questionnaire example"
    workbook.properties.description = "Wholly invented participants; no historical rows."
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "D2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, len(rows[0]) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 24 if column <= 3 else 55
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="20344D")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 36
    for row_number in range(2, len(rows) + 1):
        sheet.row_dimensions[row_number].height = 36
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    # ZIP timestamps and openpyxl's modified-at update otherwise vary on each run.
    output = BytesIO()
    with ZipFile(BytesIO(buffer.getvalue()), "r") as source, ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for filename in sorted(source.namelist()):
            payload = source.read(filename)
            if filename == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:modified\b[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>", payload,
                )
            entry = ZipInfo(filename, date_time=(2000, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.create_system = 0
            target.writestr(entry, payload)
    return output.getvalue()


def check_existing() -> int:
    failures = []
    for filename, expected in datasets().items():
        path = DATA_DIR / filename
        if not path.is_file():
            failures.append(f"{filename}: missing")
            continue
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if workbook.sheetnames != ["Sheet1"]:
                failures.append(f"{filename}: unexpected worksheets")
                continue
            actual = [list(row) for row in workbook["Sheet1"].iter_rows(values_only=True)]
            if actual != expected:
                failures.append(f"{filename}: content differs from synthetic declaration")
            if workbook.properties.creator != "Saturday matching synthetic examples":
                failures.append(f"{filename}: nonstandard author metadata")
        finally:
            workbook.close()
    if failures:
        print("\n".join(failures), file=sys.stderr)
        return 1
    print("Verified six synthetic workbooks: 8 tutors and 12 tutees per pair.")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    action = parser.add_mutually_exclusive_group()
    action.add_argument("--check", action="store_true", help="Verify existing examples without writing.")
    action.add_argument("--json", action="store_true", help="Print the synthetic declarations as JSON without writing.")
    args = parser.parse_args()
    if args.check:
        return check_existing()
    if args.json:
        print(json.dumps(datasets(), ensure_ascii=False))
        return 0
    examples = datasets()
    existing = [name for name in examples if (DATA_DIR / name).exists()]
    if existing:
        print("Refusing to overwrite existing example files. Use --check to verify them.", file=sys.stderr)
        return 2
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    # Build everything before creating the first file. Exclusive mode also guards
    # against another process creating a target between the preflight and write.
    payloads = {name: workbook_bytes(rows) for name, rows in examples.items()}
    for filename, payload in payloads.items():
        with (DATA_DIR / filename).open("xb") as target:
            target.write(payload)
    print("Created six synthetic workbooks in version_1/data.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
