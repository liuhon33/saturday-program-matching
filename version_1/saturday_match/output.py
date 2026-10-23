"""Local Excel delivery, minimal columns, exclusive creation and formula safety."""
from __future__ import annotations

import csv
from dataclasses import asdict
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def write_json(path, value):
    with Path(path).open('x', encoding='utf-8') as f:
        json.dump(value, f, indent=2, ensure_ascii=False)


def write_csv(path, headers, rows):
    with Path(path).open('x', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            # CSV is a human-review surface and must not execute user text.
            writer.writerow({k: ("'" + str(row.get(k, '')) if str(row.get(k, '')).startswith(('=', '+', '-', '@')) else row.get(k, '')) for k in headers})


def workbook(path, sheets):
    """All strings are explicit literal cells, including leading equals signs."""
    path = Path(path)
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = 'Saturday Program Matcher'
    book.properties.lastModifiedBy = 'Saturday Program Matcher'
    for title, (headers, rows) in sheets.items():
        ws = book.create_sheet(title)
        for row_number, row in enumerate([headers, *rows], 1):
            for column, value in enumerate(row, 1):
                cell = ws.cell(row_number, column)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
                if row_number == 1:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='24475C')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            width = min(52, max(12, max(len(str(c.value or '')) for c in column) + 2))
            ws.column_dimensions[column[0].column_letter].width = width
    # Opening with xb also prevents an accidental overwrite if two runs race.
    with path.open('xb') as stream:
        book.save(stream)
    book.close()


def export_cleaning(directory, loaded):
    directory = Path(directory)
    events = [e for group in loaded for e in group.events]
    ledger = [row for group in loaded for row in group.ledger]
    write_json(directory / 'cleaning_audit.json', events)
    write_json(directory / 'input_summary.json', [g.metadata for g in loaded])
    write_csv(directory / 'row_status.csv', ['role', 'source_row', 'status', 'participant_id'], ledger)
    write_csv(directory / 'corrections.template.csv', ['role', 'source_sha256', 'row', 'field', 'value'], [])
    for group in loaded:
        records = [asdict(r) for r in group.records]
        # Use JSON for typed normalized records; the file stays in private output.
        write_json(directory / f'normalized_{group.role}s.json', records)
        if group.role == 'tutor':
            headers = ['Tutor ID','Tutor Name','Canonical Subjects','Gender','Capacity','Grades','Slots','Source Row']
            rows = [[r.id,r.display_name,' | '.join(r.subjects),r.gender,r.capacity,' | '.join(r.grades),' | '.join(r.slots),r.source_row] for r in group.records]
        else:
            headers = ['Tutee ID','Tutee Name','Primary','Secondary','Tertiary','Gender Preference','Grade','Slots','Source Row']
            rows = [[r.id,r.display_name,*(list(r.preferences) + [None] * (3-len(r.preferences))),r.gender_preference,r.grade,' | '.join(r.slots),r.source_row] for r in group.records]
        workbook(directory / f'cleaned_{group.role}s.xlsx', {'Cleaned': (headers, rows)})


def export_matches(directory, result, tutees, tutors, *, include_contacts=False):
    directory = Path(directory)
    students = {s.id:s for s in tutees.records}
    teachers = {t.id:t for t in tutors.records}
    pair_headers = ['Tutee ID','Tutee Name','Tutor ID','Tutor Name','Primary Subject','Score','Urgency',
                    'Primary Hit','Secondary Hit','Tertiary Hit','Common Slots','Flags']
    pair_rows = []
    for pair in result['pairs']:
        s, t = students[pair['tutee_id']], teachers[pair['tutor_id']]
        pair_rows.append([s.id,s.display_name,t.id,t.display_name,s.preferences[0],pair['score'],pair['urgency_score'],
                          *(pair['compatibility'][k] for k in ('primary','secondary','tertiary')),
                          ' | '.join(sorted(set(s.slots) & set(t.slots))),'; '.join(pair['violations'])])
    unmatched = [[r['tutee_id'],students[r['tutee_id']].display_name,r['reason'],json.dumps(r['rejection_counts'],sort_keys=True)] for r in result['unmatched_tutees']]
    remaining = [[r['tutor_id'],teachers[r['tutor_id']].display_name,r['capacity'],r['used_capacity'],r['remaining_capacity']] for r in result['remaining_capacity']]
    sheets = {'Matches':(pair_headers,pair_rows),
              'Unmatched Tutees':(['Tutee ID','Tutee Name','Reason','Candidate Rejections'], unmatched),
              'Tutor Capacity':(['Tutor ID','Tutor Name','Capacity','Used','Remaining'], remaining)}
    if include_contacts:
        sheets['Private Contacts'] = (['Role','Participant ID','Name','Email'],
            [[g.role,identity,c['name'],c['email']] for g in (tutors,tutees) for identity,c in g.contacts.items()])
    workbook(directory / 'matches.xlsx', sheets)
    write_json(directory / 'matching_details.json', result)
