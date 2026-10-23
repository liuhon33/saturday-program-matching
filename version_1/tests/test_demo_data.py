"""Integration coverage for the public, wholly fictional form-sized examples.

Only version_1/data is read. The documented response counts and form widths are
fixed examples, not assertions made by opening any private historical inputs.
"""

from collections import Counter
from pathlib import Path
import random
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from saturday_match.config import load_config
from saturday_match.loader import load_input
from saturday_match.matcher import match_greedy


BASE = Path(__file__).resolve().parents[1]
COUNTS = {"2023": (93, 99), "2024": (80, 71), "scratch_": (91, 142)}
SHAPES = {
    "2023tutor.xlsx": {"Sheet1": (93, 27)},
    "2023tutee.xlsx": {"Sheet1": (99, 40)},
    "2024tutor.xlsx": {"Sheet1": (80, 28)},
    "2024tutee.xlsx": {"Sheet1": (71, 39)},
    "scratch_tutor.xlsx": {"Sheet1": (91, 28)},
    "scratch_tutee.xlsx": {"Sheet1": (142, 39), "Sheet2": (67, 3)},
}


class ExpandedDemoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.profiles = {name: load_config(BASE / "config" / f"{name}.yaml") for name in ("default", "synthetic")}
        cls.loaded = {}
        for name, config in cls.profiles.items():
            for prefix in COUNTS:
                for role in ("tutor", "tutee"):
                    cls.loaded[name, prefix, role] = load_input(BASE / "data" / f"{prefix}{role}.xlsx", role, config)

    def test_documented_form_shapes_and_supplementary_sheet(self):
        for filename, sheets in SHAPES.items():
            with self.subTest(filename=filename):
                book = load_workbook(BASE / "data" / filename, read_only=True)
                try:
                    self.assertEqual(book.sheetnames, list(sheets))
                    for name, (responses, columns) in sheets.items():
                        rows = list(book[name].iter_rows(values_only=True))
                        self.assertEqual(len(rows) - 1, responses)
                        self.assertTrue(all(len(row) == columns for row in rows))
                finally:
                    book.close()

    def test_all_primary_response_rows_are_valid_in_both_profiles(self):
        for name in self.profiles:
            for prefix, counts in COUNTS.items():
                for role, expected in zip(("tutor", "tutee"), counts):
                    with self.subTest(profile=name, pair=prefix, role=role):
                        loaded = self.loaded[name, prefix, role]
                        self.assertFalse(loaded.errors)
                        self.assertEqual(len(loaded.records), expected)
                        self.assertEqual(len(loaded.ledger), expected)
                        self.assertEqual(loaded.metadata["invalid_rows"], 0)
                        self.assertEqual(loaded.metadata["blank_rows"], 0)
                        self.assertTrue(all(row["status"] == "valid" for row in loaded.ledger))

    def test_same_names_remain_distinct_and_every_tutee_is_accounted_for(self):
        for name, config in self.profiles.items():
            for prefix, (tutor_count, tutee_count) in COUNTS.items():
                with self.subTest(profile=name, pair=prefix):
                    tutors = self.loaded[name, prefix, "tutor"]
                    tutees = self.loaded[name, prefix, "tutee"]
                    for loaded in (tutors, tutees):
                        self.assertTrue(any(count > 1 for count in Counter(p.display_name for p in loaded.records).values()))
                        self.assertEqual(len({person.id for person in loaded.records}), len(loaded.records))
                        self.assertIn("same_name_distinct_ids_retained", {event["rule"] for event in loaded.events})
                    result = match_greedy(tutees.records, tutors.records, config)
                    summary = result["summary"]
                    self.assertEqual(summary["tutor_count"], tutor_count)
                    self.assertEqual(summary["tutee_count"], tutee_count)
                    self.assertEqual(summary["matched_count"] + summary["unmatched_count"], tutee_count)
                    self.assertEqual(summary["hard_constraint_violations"], 0)
                    self.assertTrue(summary["conservation_passed"])
                    accounted = [pair["tutee_id"] for pair in result["pairs"]] + [row["tutee_id"] for row in result["unmatched_tutees"]]
                    self.assertCountEqual(accounted, [person.id for person in tutees.records])
                    self.assertEqual(result, match_greedy(tutees.records, tutors.records, config))

    def test_reordering_form_responses_keeps_identities_and_matches(self):
        config = self.profiles["default"]
        with tempfile.TemporaryDirectory(prefix="saturday-demo-order-") as directory:
            for prefix in COUNTS:
                reordered = {}
                for role in ("tutor", "tutee"):
                    source = load_workbook(BASE / "data" / f"{prefix}{role}.xlsx", read_only=True)
                    try:
                        rows = list(source["Sheet1"].iter_rows(values_only=True))
                    finally:
                        source.close()
                    header, responses = rows[0], rows[1:]
                    random.Random(2024).shuffle(responses)
                    target = Path(directory) / f"{prefix}{role}.xlsx"
                    book = Workbook()
                    sheet = book.active
                    sheet.title = "Sheet1"
                    for row in [header, *responses]:
                        sheet.append(row)
                    book.save(target)
                    book.close()
                    reordered[role] = load_input(target, role, config)
                    self.assertFalse(reordered[role].errors)
                    self.assertEqual({p.id for p in reordered[role].records}, {p.id for p in self.loaded["default", prefix, role].records})
                with self.subTest(pair=prefix):
                    original = match_greedy(self.loaded["default", prefix, "tutee"].records, self.loaded["default", prefix, "tutor"].records, config)
                    shuffled = match_greedy(reordered["tutee"].records, reordered["tutor"].records, config)
                    self.assertEqual(shuffled, original)


if __name__ == "__main__":
    unittest.main()
